"""Extracts 'O·O8 Condition Modifiers' into onboarding_o8.json.

Source: v39sEng2.xlsx, sheet 'O·O8 Condition Modifiers'.
'01_IMPORT_MANIFEST' order 78, role ONBOARDING, import YES, backend Yes.

    python -m sahacore.data.build_onboarding_o8 <path-to-workbook>

The authority for ONB-008. Onboarding step 5, feeding Layers A and C. Unlike
every other O-sheet this one has NO numbered equations -- it is a single
ten-row table of conditions, plus one rule stated in prose above it.

THE RULE IS THE MOST IMPORTANT THING ON THE SHEET, and it is a safety rule:

    "COMPATIBILITY RULE: any legacy positive F_bio multiplier m is
     interpreted as an ODDS multiplier and converted to Delta_logit_abs =
     ln(m). The production equation is
     F_abs = F_max * sigmoid(logit(F_base/F_max) + sum(Delta_logit_abs));
     no direct multiplication may exceed [0,1]."

An absorbed fraction is bounded. Multiplying one by 1.5 is not, so the sheet
forbids doing it directly and gives the transformation that keeps the result
in range: work in log-odds, sum there, and come back through a sigmoid.
`F_max` is a real column in the 81-nutrient registry, so this composes with
what is already imported. It is transcribed verbatim and implemented in
sahacore/onboarding/condition_modifiers.py.

FIVE OF THE TEN ROWS ARE GATED -- FOUR OF THEM GATE A NUMBER. The table does not
only say "eta_Z7 x1.5" -- it says "eta_Z7 x1.5 ONLY WHEN CALIBRATED", "eta_Z3
x1.2 only if symptoms support it", "eta_Z11 x1.5 only when confirmed". A
reader who took the number and dropped the clause would apply a 50% damage
sensitivity increase the sheet explicitly withheld. So the gate text is
extracted as its own field and a modifier that carries one is marked
conditional; what "calibrated" or "confirmed" MEANS is defined nowhere, which
is recorded rather than guessed.

The Evidence role column is doing the same work in the other direction --
"Safety/target modifier; do not force K malabsorption", "Timing, not global
absorption extent", "Clinical context; not an absorption multiplier". Those
are warnings against a specific misreading, and they are kept verbatim.

TWO GAPS, both checked in code:

  The Z-pathways column names pathways no modifier touches. Hypertension
  declares "Z7, Z9" and modifies only Z7; CKD declares "Z9, Z14" and modifies
  only Z9. A pathway declared affected with no modifier is a claim with no
  arithmetic behind it.

  The interface offers a condition this sheet has no row for. Step 5's
  gastrointestinal question names "IBS, GERD, Celiac, UC, NAFLD" and there is
  no UC row -- the same shape as O7's Intermittent Fasting. And of Step 5's
  five condition questions only two name any condition at all, one of those
  trailing off with "etc.", so exactly one gives a checkable list. The full
  condition set is unknown, and O8's ten rows cannot be checked for
  completeness in either direction.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

from sahacore.data.sheet_header import check_header

SHEET = "O·O8 Condition Modifiers"
FIRST_COL = 3
DATA_DIR = Path(__file__).parent
OUT = DATA_DIR / "onboarding_o8.json"

# The two prose rules sit in column B, left of the table.
SCOPE_ROW = 5
COMPATIBILITY_ROW = 7
RULE_COL = 2

HEADER_ROW = 9
DATA_ROWS = range(10, 20)
LABELS = ["Condition", "eta / state modifier", "Bounded absorption effect",
          "Target / reference adjustment", "Z-pathways", "Evidence role"]

EXPECTED_CONDITIONS = 10

# 'eta_Z1 ×1.5; eta_Z6 ×2.0'. The sheet uses the multiplication sign, not x.
_MODIFIER = re.compile(r"eta_(Z\d+)\s*[×x*]\s*(\d+(?:\.\d+)?)")

# 'Z7, Z9' in the Z-pathways column.
_PATHWAY = re.compile(r"\bZ\d+\b")

# What turns a stated modifier into a conditional one. Kept as a tuple of
# openers rather than a single regex so a new one fails loudly instead of
# being half-matched.
_GATE_OPENERS = ("only if", "only when", "when ")


def _text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _cell(ws, row: int, offset: int):
    return ws.cell(row=row, column=FIRST_COL + offset).value


def _gate(modifier_text: str) -> str | None:
    """The clause that withholds a modifier until something is true."""
    lowered = modifier_text.lower()
    for opener in _GATE_OPENERS:
        index = lowered.find(opener)
        if index != -1:
            return modifier_text[index:].strip()
    return None


def extract(path: str) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET]
    check_header(ws, SHEET, HEADER_ROW, FIRST_COL, LABELS)

    scope = _text(ws.cell(row=SCOPE_ROW, column=RULE_COL).value)
    compatibility = _text(ws.cell(row=COMPATIBILITY_ROW, column=RULE_COL).value)

    conditions = []
    for row in DATA_ROWS:
        condition = _text(_cell(ws, row, 0))
        if condition is None:
            raise SystemExit(f"{SHEET}: row {row} has no condition.")
        modifier_text = _text(_cell(ws, row, 1)) or ""
        pathways_text = _text(_cell(ws, row, 4)) or ""

        modifiers = [{"z_pathway": pathway, "factor": float(factor)}
                     for pathway, factor in _MODIFIER.findall(modifier_text)]
        declared = _PATHWAY.findall(pathways_text)
        modified = {m["z_pathway"] for m in modifiers}

        conditions.append({
            "condition": condition,
            "source_row": row,
            "modifier_text": modifier_text,
            "modifiers": modifiers,
            # The clause that withholds the modifier. Extracted as its own
            # field because dropping it would apply a factor the sheet
            # deliberately did not.
            "gate": _gate(modifier_text),
            "bounded_absorption_effect": _text(_cell(ws, row, 2)),
            "target_adjustment": _text(_cell(ws, row, 3)),
            "z_pathways_text": pathways_text,
            "z_pathways_declared": declared,
            # Declared affected, with no arithmetic saying how.
            "z_pathways_without_a_modifier": [p for p in declared
                                              if p not in modified],
            "evidence_role": _text(_cell(ws, row, 5)),
        })

    wb.close()
    return {"sheet": SHEET, "scope": scope,
            "compatibility_rule": compatibility, "conditions": conditions}


def check(data: dict) -> None:
    conditions = data["conditions"]

    if len(conditions) != EXPECTED_CONDITIONS:
        raise SystemExit(f"{SHEET}: expected {EXPECTED_CONDITIONS} conditions, "
                         f"got {len(conditions)}.")
    if len({c["condition"] for c in conditions}) != EXPECTED_CONDITIONS:
        raise SystemExit(f"{SHEET}: the condition names are not distinct.")

    for condition in conditions:
        for field in ("modifier_text", "bounded_absorption_effect",
                      "target_adjustment", "z_pathways_text", "evidence_role"):
            if not condition[field]:
                raise SystemExit(
                    f"{SHEET}: {condition['condition']!r} has no {field}. The "
                    "evidence role in particular is a warning against a "
                    "misreading and must not go missing.")
        if not condition["z_pathways_declared"]:
            raise SystemExit(
                f"{SHEET}: {condition['condition']!r} names no Z-pathway: "
                f"{condition['z_pathways_text']!r}")

    # THE COMPATIBILITY RULE, which is the safety-critical part of this sheet.
    rule = data["compatibility_rule"] or ""
    for fragment in ("odds multiplier", "logit", "sigmoid", "F_max",
                     "[0,1]"):
        if fragment.lower() not in rule.lower():
            raise SystemExit(
                f"{SHEET}: the compatibility rule no longer mentions "
                f"{fragment!r}: {rule!r}. It is what stops a 1.5x multiplier "
                "being applied directly to a bounded absorbed fraction, and "
                "sahacore.onboarding.condition_modifiers implements it "
                "verbatim.")

    # THE GATES. Five rows withhold their modifier until something is true,
    # and four of those five withhold a NUMBER. The distinction matters:
    # Celiac's gate qualifies a described modifier ("use inflammation/repair
    # modifier when active") with no factor attached, while the other four
    # each hold back a specific multiplier that a careless reader would
    # apply unconditionally.
    gated = [c for c in conditions if c["gate"]]
    gated_numbers = [c for c in gated if c["modifiers"]]
    if len(gated) != 5 or len(gated_numbers) != 4:
        raise SystemExit(
            f"{SHEET}: {len(gated)} rows carry a gate and {len(gated_numbers)} "
            "of them gate a factor; the finding was recorded against 5 and 4. "
            f"Gated: {[c['condition'] for c in gated]}")
    for condition in gated_numbers:
        # A gate must come AFTER the factor it withholds; a leading one would
        # mean the sheet is qualifying something else.
        if condition["modifier_text"].index(condition["gate"]) == 0:
            raise SystemExit(
                f"{SHEET}: {condition['condition']!r} has its gate before its "
                f"factor: {condition['modifier_text']!r}. Re-read it -- the "
                "gate must qualify the number, not the other way round.")

    # A pathway declared affected with no modifier saying how.
    unmodified = {c["condition"]: c["z_pathways_without_a_modifier"]
                  for c in conditions if c["z_pathways_without_a_modifier"]}
    if not unmodified:
        raise SystemExit(
            f"{SHEET}: every declared Z-pathway now carries a modifier. The "
            "finding that several do not is resolved -- re-read "
            "docs/parameter-gaps.md and replace this check.")

    # THE INTERFACE OFFERS A CONDITION THIS SHEET HAS NO ROW FOR.
    ui = json.loads((DATA_DIR / "step_questions.json").read_text(encoding="utf-8"))
    step_5 = [q for q in ui["questions"] if q["step_number"] == 5]
    if not step_5:
        raise SystemExit(f"{SHEET}: the UI contract has no Step 5 questions.")
    # Two of the five name any condition at all, and one of those two trails
    # off with "etc." -- so exactly ONE question in Step 5 gives a list that
    # can be checked against this sheet. The other four leave the condition
    # set unknown, which is why O8's ten rows cannot be checked for
    # completeness in either direction.
    naming = [q for q in step_5 if ":" in q["answer_options_text"]]
    complete = [q for q in naming if "etc" not in q["answer_options_text"].lower()]
    if (len(naming), len(complete)) != (2, 1):
        raise SystemExit(
            f"{SHEET}: {len(naming)} of Step 5's {len(step_5)} condition "
            f"questions name a condition and {len(complete)} of those give a "
            "list that does not trail off; the finding was recorded against "
            "2 and 1. If more now enumerate, check the condition list "
            "properly rather than keeping this note.")
    named = complete[0]["answer_options_text"]
    if "UC" not in named:
        raise SystemExit(
            f"{SHEET}: Step 5 no longer names UC: {named!r}. It was the "
            "condition the interface offers and this sheet has no row for.")
    if any("UC" == c["condition"] or "olitis" in c["condition"]
           for c in conditions):
        raise SystemExit(
            f"{SHEET}: a UC row now exists, so the finding is resolved. "
            "Re-read docs/parameter-gaps.md.")


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("usage: python -m sahacore.data.build_onboarding_o8 "
                         "<path-to-workbook>")
    data = extract(sys.argv[1])
    check(data)
    OUT.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n",
                   encoding="utf-8")

    conditions = data["conditions"]
    modifiers = sum(len(c["modifiers"]) for c in conditions)
    gated = [c["condition"] for c in conditions if c["gate"]]
    unmodified = sorted({p for c in conditions
                         for p in c["z_pathways_without_a_modifier"]})
    print(f"wrote {OUT.name}")
    print(f"  {len(conditions)} conditions, {modifiers} eta modifiers")
    numeric = [c["condition"] for c in conditions if c["gate"] and c["modifiers"]]
    print(f"  {len(gated)} rows GATED, {len(numeric)} of them gating a "
          f"factor: {', '.join(numeric)}")
    print(f"  {len(unmodified)} Z-pathways declared affected with no modifier: "
          + ", ".join(unmodified))
    print("  compatibility rule captured: F_abs = F_max*sigmoid(logit(...) + "
          "sum(dlogit)), no direct multiplication")


if __name__ == "__main__":
    main()

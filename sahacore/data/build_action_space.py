"""Extracts the Layer H action space into action_space_127.json.

Source: v39sEng2.xlsx, sheet 'Action_Space' -- "PHASE 1: Action Space
Enumeration — 127 Discrete Bandit Actions". Header at row 9, actions at rows
10-136.

    python -m sahacore.data.build_action_space <path-to-workbook>

THE SHEET HOLDS THREE TABLES, AND THEY ARE NOT THE SAME THING.

  actions (127)    the outcome-bearing arms the LinUCB bandit ranks:
                   81 nutrient (41 INCREASE + 40 DECREASE) + 20 activity +
                   15 sleep/stress + 11 timing.
  info (7)         INFO-01..INFO-07. The sheet is explicit that these are
                   "NOT PART OF THE 127 OUTCOME-BEARING ACTIONS" and run
                   under a separate value-of-information policy. Folding them
                   into the action table would make the bandit rank a request
                   to resync a wearable against a request to change protein
                   intake.
  phases (4)       the phased arm-activation schedule: which arms are
                   eligible at all, given how much data a user has.

WHY SAFETY IS NOT A COLUMN HERE. The sheet carries a 'VETO Rule Cross-Ref'
cell per action, and it is prose -- "VETO: warfarin (rule #1)", "Caution:
oxalate nephropathy at >2g/d", "None at dietary doses". It is transcribed as
written and NOT parsed into foreign keys, because the sheet says where the
gate actually comes from:

    "★ Safety: Layer H builds the action × active-rule gate from the
     versioned VETO registry at load time. Do not hard-code."

So the authority is engine_internal.veto_drug_nutrient, and this column is a
human-readable note beside it.

ACTION 127 IS HELD, DELIBERATELY. The phase table's own Phase 3 cell reads:

    "126 / 127 — full ACTIVATABLE set, NOT the full ontology. Action 127
     (INCREASE Dietary Nitrate, nutrient 81) stays inactive pending
     VN-01…VN-07 clinical sign-off AND regeneration of the 127×140
     action×rule gate. This is a deliberate safety hold, not an incomplete
     rollout. Do NOT write an acceptance test asserting 127 active arms at
     Phase 3, and do NOT activate arm 127 to make the count 'full'."

That instruction is followed exactly. The ontology keeps all 127 rows;
`activation_hold` carries that sentence, verbatim from that cell, on action
127 and on no other row -- the builder raises if it would land anywhere else.
"""
import json
import sys
from pathlib import Path

import openpyxl

from sahacore.data.sheet_header import check_header

SHEET = "Action_Space"
HEADER_ROW = 9
FIRST_ACTION_ROW = 10
OUT = Path(__file__).parent / "action_space_127.json"

# Column B onward; column A is the sheet's indent margin and is always empty.
COLUMNS = [
    "action_id", "category", "subcategory", "action", "magnitude_bin",
    "primary_target", "cluster_affected", "veto_cross_ref",
    "default_reward_prior", "policy_class", "online_exploration_rule",
    "physiological_uncertainty_rule", "required_evaluation",
]
HEADER_LABELS = [
    "Action ID", "Category", "Subcategory", "Action", "Magnitude Bin",
    "Primary Nutrient/Target", "Cluster Affected", "VETO Rule Cross-Ref",
    "Default Reward Prior", "Policy Class", "Online Exploration Rule",
    "Physiological Uncertainty Rule", "Required Evaluation",
]

INFO_HEADER_ROW = 201
# The information actions carry the same policy columns as the 127 main arms
# -- VOI score, the state they affect, their safety position, how long the
# answer stays useful, policy class, exploration rule, uncertainty rule and
# the evaluation each needs. An earlier version of this extractor took the
# first five and dropped these eight, including Safety/VETO. The header check
# in sahacore.data.sheet_header now refuses a header wider than the labels
# given, which is what surfaced them.
INFO_COLUMNS = ["info_id", "information_action", "expected_information_gain",
                "user_burden", "eligibility", "voi_score",
                "state_or_output_affected", "safety_veto", "expiration",
                "policy_class", "exploration", "state_uncertainty",
                "evaluation"]
INFO_HEADER_LABELS = ["Info ID", "Information action",
                      "Expected information gain target", "User burden",
                      "Eligibility", "VOI score", "State/output affected",
                      "Safety/VETO", "Expiration", "Policy class",
                      "Exploration", "State uncertainty", "Evaluation"]

PHASE_HEADER_ROW = 155
# Three more columns the same partial-header check hid: how many samples a
# phase needs before it may advance, what convergence means for it, and its
# safety notes. A rollout schedule without its thresholds is a list of names.
PHASE_COLUMNS = ["phase", "trigger_condition", "new_arms_activated",
                 "cumulative_arms", "active_categories", "sample_threshold",
                 "convergence_criterion", "safety_notes"]
PHASE_HEADER_LABELS = ["Phase", "Trigger Condition", "New Arms Activated",
                       "Cumulative Arms", "Active Categories",
                       "Sample Threshold", "Convergence Criterion",
                       "Safety Notes"]

# The one action the sheet holds inactive, and the cell that says so. Both are
# named here so the hold cannot silently move to another arm or be dropped.
HELD_ACTION_ID = 127
HOLD_SOURCE_CELL = (159, 5)          # phase table, Phase 3, 'Cumulative Arms'

# The sheet's own breakdown line, checked against what was extracted.
DECLARED_CATEGORY_COUNTS = {
    "Nutrient": 81, "Activity": 20, "Sleep/Stress": 15, "Timing": 11,
}
DECLARED_NUTRIENT_VERBS = {"INCREASE": 41, "DECREASE": 40}


def _cell(ws, row: int, col: int) -> str | None:
    value = ws.cell(row=row, column=col).value
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _table(ws, header_row: int, columns: list[str], labels: list[str],
           what: str, stop_after: int | None = None) -> list[dict]:
    check_header(ws, f"{SHEET} ({what})", header_row, 2, labels)
    rows = []
    last = stop_after or ws.max_row
    for r in range(header_row + 1, last + 1):
        if _cell(ws, r, 2) is None:
            continue
        row = {"source_row": r}
        for i, name in enumerate(columns, start=2):
            row[name] = _cell(ws, r, i)
        rows.append(row)
    return rows


def extract(workbook_path: str) -> dict:
    ws = openpyxl.load_workbook(workbook_path, data_only=True)[SHEET]

    actions = []
    check_header(ws, f"{SHEET} (action)", HEADER_ROW, 2, HEADER_LABELS)
    for r in range(FIRST_ACTION_ROW, ws.max_row + 1):
        raw = _cell(ws, r, 2)
        if raw is None:
            continue
        try:
            action_id = int(raw)
        except ValueError:
            break                    # the first banner after the action block
        row = {"source_row": r, "action_id": action_id}
        for i, name in enumerate(COLUMNS[1:], start=3):
            row[name] = _cell(ws, r, i)
        row["default_reward_prior"] = float(row["default_reward_prior"])
        row["activation_hold"] = None
        actions.append(row)

    hold = _cell(ws, *HOLD_SOURCE_CELL)
    held = [a for a in actions if a["action_id"] == HELD_ACTION_ID]
    if len(held) != 1:
        raise SystemExit(f"action {HELD_ACTION_ID} is not in the sheet exactly once")
    if f"Action {HELD_ACTION_ID}" not in (hold or ""):
        raise SystemExit(
            f"cell {HOLD_SOURCE_CELL} no longer names action {HELD_ACTION_ID}; "
            "the safety hold must not be attached to an arm on a guess")
    held[0]["activation_hold"] = hold

    return {
        "actions": actions,
        "info_actions": _table(ws, INFO_HEADER_ROW, INFO_COLUMNS,
                               INFO_HEADER_LABELS, "info"),
        "phases": _table(ws, PHASE_HEADER_ROW, PHASE_COLUMNS,
                         PHASE_HEADER_LABELS, "phase", stop_after=159),
    }


def check(data: dict) -> None:
    actions = data["actions"]
    ids = [a["action_id"] for a in actions]
    if sorted(ids) != list(range(1, 128)):
        raise SystemExit(f"expected action ids 1..127, got {len(ids)} ids")

    counts: dict[str, int] = {}
    for a in actions:
        counts[a["category"]] = counts.get(a["category"], 0) + 1
    if counts != DECLARED_CATEGORY_COUNTS:
        raise SystemExit(
            f"category counts {counts} do not match the sheet's own breakdown "
            f"line, which states {DECLARED_CATEGORY_COUNTS}")

    verbs: dict[str, int] = {}
    for a in actions:
        if a["category"] == "Nutrient":
            verb = a["action"].split()[0]
            verbs[verb] = verbs.get(verb, 0) + 1
    if verbs != DECLARED_NUTRIENT_VERBS:
        raise SystemExit(f"nutrient verbs {verbs} != {DECLARED_NUTRIENT_VERBS}")

    on_hold = [a["action_id"] for a in actions if a["activation_hold"]]
    if on_hold != [HELD_ACTION_ID]:
        raise SystemExit(f"activation_hold landed on {on_hold}, not [{HELD_ACTION_ID}]")

    if len(data["info_actions"]) != 7:
        raise SystemExit(f"expected 7 INFO actions, got {len(data['info_actions'])}")
    if [i["info_id"] for i in data["info_actions"]] != [f"INFO-{n:02d}" for n in range(1, 8)]:
        raise SystemExit("INFO ids are not the contiguous INFO-01..INFO-07 range")
    if len(data["phases"]) != 4:
        raise SystemExit(f"expected 4 activation phases, got {len(data['phases'])}")


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("usage: python -m sahacore.data.build_action_space <workbook.xlsx>")

    data = extract(sys.argv[1])
    check(data)
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    print(f"wrote the action space to {OUT.name}")
    print(f"   outcome-bearing actions : {len(data['actions'])}  {DECLARED_CATEGORY_COUNTS}")
    print(f"   held inactive           : "
          f"{[a['action_id'] for a in data['actions'] if a['activation_hold']]}")
    print(f"   information actions     : {len(data['info_actions'])}  (separate VoI policy)")
    print(f"   activation phases       : {len(data['phases'])}")
    for p in data["phases"]:
        print(f"     {p['phase'].splitlines()[0]:<9} {p['cumulative_arms'][:56]}")


if __name__ == "__main__":
    main()

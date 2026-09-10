"""Extracts 'O·O9 Drug-Nutrient Mods' into onboarding_o9.json.

Source: v39sEng2.xlsx, sheet 'O·O9 Drug-Nutrient Mods'.
'01_IMPORT_MANIFEST' order 79, role ONBOARDING, import YES, backend Yes.

    python -m sahacore.data.build_onboarding_o9 <path-to-workbook>

The authority for ONB-009. Onboarding step 7, twenty drug-nutrient rows, each
with its mechanism, its severity and where it lands in the engine.

O8'S RULE, APPLIED AND CHECKABLE. Five rows carry BOTH a legacy multiplier
and its converted log-odds shift, and the conversion is exactly the rule
'O·O8 Condition Modifiers' states one sheet earlier:

    Metformin  x B12   m = 0.70   Delta_logit_abs = -0.356675 = ln(0.70)
    PPIs       x Mg    m = 0.75                    -0.287682 = ln(0.75)
    PPIs       x Ca    m = 0.80                    -0.223144 = ln(0.80)
    PPIs       x B12   m = 0.85                    -0.162519 = ln(0.85)
    PPIs       x Fe    m = 0.80                    -0.223144 = ln(0.80)

Every one is correct to six decimals. This extractor recomputes all five, so
the two sheets are held to each other rather than each taken on its own word.
It is the first place in the build where one sheet's rule is verified by
another sheet's arithmetic.

FIFTEEN ROWS CARRY NO NUMBER, AND THAT IS DELIBERATE. Their 'Legacy m / rule'
cell holds an action class instead: TIMING, MONITOR, VETO, VETO/MONITOR,
CLEARANCE, or N/A. Those are not absorption multipliers and must never be
treated as one -- which the sheet says explicitly, twice, in the strongest
terms available to it:

    Levothyroxine    "Layer H: timing VETO; do not alter nutrient F_abs"
    Fluoroquinolones "Layer H: timing VETO; nutrient F_abs unchanged"

In both, chelation reduces absorption of the DRUG, not of the nutrient. Code
that inverted that would cut a nutrient target because the user takes a
thyroid tablet. The production target is kept verbatim for exactly that
reason, and the extractor refuses a row whose rule is TIMING and whose
Delta_logit_abs is not N/A.

Statins x CoQ10 makes the same distinction the other way: "Biosynthesis
depletion; not intestinal F_abs" -- a real CRITICAL effect that is simply not
an absorption one.

WHAT THE INTERFACE COLLECTS AND THIS SHEET DOES NOT MODEL. Step 7 names 22
medications across five category rows. Three are the same drug spelled
differently -- "ACE Inhibitors & ARBs" against "ACE inhibitors/ARBs",
"Insulin & Sulfonylureas" against "Insulin/sulfonylureas", "Oral
Contraceptives" against "Oral contraceptives" -- and normalising case and the
choice of '&' or '/' bridges those without changing a word. Eight remain:
seven with no row of any kind, plus "Diuretics", which is a BROADER class
than the sheet's "Thiazide/loop diuretics" and so a different question.
And one drug goes the other way -- Warfarin is modelled, with a CRITICAL
Vitamin K veto, and is not among the 22 the interface names. Step 7's first
row is a free-text "Search bar + categories", so the named ones are examples
rather than the whole list; both directions are recorded rather than resolved.
"""
import json
import math
import re
import sys
from pathlib import Path

import openpyxl

from sahacore.data.sheet_header import check_header

SHEET = "O·O9 Drug-Nutrient Mods"
FIRST_COL = 3
DATA_DIR = Path(__file__).parent
OUT = DATA_DIR / "onboarding_o9.json"

SCOPE_ROW = 5
BANNER_ROW = 7
RULE_COL = 2

HEADER_ROW = 9
DATA_ROWS = range(10, 30)
LABELS = ["Drug / class", "Nutrient or substrate", "Legacy m / rule",
          "Δlogit_abs", "Mechanism class", "Severity", "Production target"]

EXPECTED_ROWS = 20

# The action classes a row may carry instead of a multiplier. Not absorption
# effects: a row marked TIMING changes when something is taken, and a row
# marked MONITOR asks for a measurement.
ACTION_CLASSES = {"TIMING", "MONITOR", "VETO", "VETO/MONITOR", "CLEARANCE",
                  "N/A"}

SEVERITIES = {"CRITICAL", "MODERATE", "LOW"}

# WHERE THIS SHEET AND THE 339-ROW VETO REGISTRY COVER THE SAME HAZARD.
#
# Bridged BY HAND and kept to what has actually been read, rather than
# matched on drug names. A substring match on "Metformin" pulls in seven VETO
# rows about different nutrients, and comparing severities across those would
# manufacture disagreements that are not there.
#
# So this holds one pair, verified row by row: the VETO rules are
# VETO-DN-0265 and VETO-DN-0267, both CRITICAL, both actioned "STABLE PATTERN
# - discuss with prescriber". O9 rates the same drug class MODERATE and
# actions it MONITOR. A prescriber referral against a measurement, for a
# hypoglycaemia hazard.
VETO_COMPARISONS = {
    "Insulin/sulfonylureas": {
        "veto_rule_ids": ("VETO-DN-0265", "VETO-DN-0267"),
        "veto_nutrient": "Carbohydrate intake",
        "veto_severity": "CRITICAL",
    },
}

# The sheet rounds its converted shifts to six decimals.
LOG_TOLERANCE = 5e-7

_NUMBER = re.compile(r"^-?\d+(?:\.\d+)?$")

# The interface and this sheet write the same drug class differently:
# "ACE Inhibitors & ARBs" against "ACE inhibitors/ARBs", "Insulin &
# Sulfonylureas" against "Insulin/sulfonylureas", "Oral Contraceptives"
# against "Oral contraceptives" -- case and the choice of '&' or '/'.
#
# Normalising only THOSE is a narrow, defensible bridge: it changes no word.
# It deliberately does NOT match "Diuretics" to "Thiazide/loop diuretics",
# which is a broader class against a narrower one and a real question.
_SEPARATOR = re.compile(r"\s*[&/]\s*")
_SPACES = re.compile(r"\s+")


def _normalise(name: str) -> str:
    return _SPACES.sub(" ", _SEPARATOR.sub("|", name.strip().lower()))


def _text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _cell(ws, row: int, offset: int):
    return ws.cell(row=row, column=FIRST_COL + offset).value


def _maybe_number(text: str | None) -> float | None:
    if text is None or not _NUMBER.match(text):
        return None
    return float(text)


def extract(path: str) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET]
    check_header(ws, SHEET, HEADER_ROW, FIRST_COL, LABELS)

    scope = _text(ws.cell(row=SCOPE_ROW, column=RULE_COL).value)
    banner = _text(ws.cell(row=BANNER_ROW, column=RULE_COL).value)

    interactions = []
    for row in DATA_ROWS:
        drug = _text(_cell(ws, row, 0))
        if drug is None:
            raise SystemExit(f"{SHEET}: row {row} has no drug.")
        rule_text = _text(_cell(ws, row, 2))
        shift_text = _text(_cell(ws, row, 3))

        interactions.append({
            "drug": drug,
            "nutrient": _text(_cell(ws, row, 1)),
            "source_row": row,
            "rule_text": rule_text,
            # Set only on the five rows that give a multiplier; the rest hold
            # an action class, which is not an absorption effect.
            "legacy_multiplier": _maybe_number(rule_text),
            "delta_logit_abs": _maybe_number(shift_text),
            "delta_logit_text": shift_text,
            "mechanism_class": _text(_cell(ws, row, 4)),
            "severity": _text(_cell(ws, row, 5)),
            "production_target": _text(_cell(ws, row, 6)),
        })

    wb.close()
    return {"sheet": SHEET, "scope": scope, "banner": banner,
            "interactions": interactions,
            "severity_scale": sorted({r["severity"] for r in interactions}),
            "severity_disagreements": _veto_disagreements(interactions)}


def _veto_disagreements(interactions: list[dict]) -> list[dict]:
    """This sheet against the drug-nutrient VETO registry, for the pairs that
    have been read against each other by hand."""
    veto = json.loads(
        (DATA_DIR / "veto_drug_nutrient_339.json").read_text(encoding="utf-8"))
    rows = veto if isinstance(veto, list) else next(
        v for v in veto.values() if isinstance(v, list))
    by_id = {r["rule_id"]: r for r in rows}

    disagreements = []
    for drug, expected in VETO_COMPARISONS.items():
        here = next((r for r in interactions if r["drug"] == drug), None)
        if here is None:
            raise SystemExit(
                f"{SHEET}: {drug!r} is bridged to the VETO registry and has "
                "no row here any more.")
        for rule_id in expected["veto_rule_ids"]:
            if rule_id not in by_id:
                raise SystemExit(
                    f"{SHEET}: {rule_id} is no longer in the VETO registry; "
                    "the comparison was recorded against it.")
            veto_row = by_id[rule_id]
            if veto_row["severity"] != expected["veto_severity"]:
                raise SystemExit(
                    f"{SHEET}: {rule_id} is now {veto_row['severity']!r}, not "
                    f"{expected['veto_severity']!r}. The disagreement with "
                    "this sheet may be resolved -- re-read both.")
            if veto_row["severity"] != here["severity"]:
                disagreements.append({
                    "drug": drug,
                    "o9_nutrient": here["nutrient"],
                    "o9_severity": here["severity"],
                    "o9_rule": here["rule_text"],
                    "o9_production_target": here["production_target"],
                    "veto_rule_id": rule_id,
                    "veto_nutrient": veto_row["nutrient_or_food"],
                    "veto_severity": veto_row["severity"],
                    "veto_action": veto_row["action"],
                })
    return disagreements


def _ui_medications() -> set[str]:
    data = json.loads((DATA_DIR / "step_questions.json").read_text(encoding="utf-8"))
    named: set[str] = set()
    for question in data["questions"]:
        if question["step_number"] == 7 and question["is_choice"]:
            named.update(question["answer_options"])
    return named


def check(data: dict) -> None:
    interactions = data["interactions"]

    if len(interactions) != EXPECTED_ROWS:
        raise SystemExit(f"{SHEET}: expected {EXPECTED_ROWS} rows, got "
                         f"{len(interactions)}.")
    for row in interactions:
        for field in ("drug", "nutrient", "rule_text", "delta_logit_text",
                      "mechanism_class", "severity", "production_target"):
            if not row[field]:
                raise SystemExit(
                    f"{SHEET}: row {row['source_row']} has no {field}. The "
                    "production target in particular says whether an effect "
                    "touches nutrient absorption at all.")
        if row["severity"] not in SEVERITIES:
            raise SystemExit(
                f"{SHEET}: row {row['source_row']} has severity "
                f"{row['severity']!r}, not one of {sorted(SEVERITIES)}.")

    # O8'S RULE, RECOMPUTED. This is the whole point of importing O8 first.
    quantified = [r for r in interactions if r["legacy_multiplier"] is not None]
    if len(quantified) != 5:
        raise SystemExit(
            f"{SHEET}: {len(quantified)} rows give a legacy multiplier and the "
            "finding was recorded against five.")
    for row in quantified:
        multiplier = row["legacy_multiplier"]
        if not 0 < multiplier <= 1:
            raise SystemExit(
                f"{SHEET}: row {row['source_row']} has m = {multiplier}. Every "
                "quantified interaction on this sheet REDUCES absorption; a "
                "multiplier above 1 would need reading, not importing.")
        if row["delta_logit_abs"] is None:
            raise SystemExit(
                f"{SHEET}: row {row['source_row']} gives m = {multiplier} and "
                f"no Delta_logit_abs: {row['delta_logit_text']!r}")
        computed = math.log(multiplier)
        if abs(computed - row["delta_logit_abs"]) > LOG_TOLERANCE:
            raise SystemExit(
                f"{SHEET}: row {row['source_row']} states "
                f"Delta_logit_abs = {row['delta_logit_abs']} for m = "
                f"{multiplier}, but ln(m) = {computed:.9f}. O·O8's "
                "compatibility rule says these must be the same number.")

    # THE FIFTEEN ROWS THAT ARE NOT ABSORPTION EFFECTS.
    for row in interactions:
        if row["legacy_multiplier"] is not None:
            continue
        if row["rule_text"] not in ACTION_CLASSES:
            raise SystemExit(
                f"{SHEET}: row {row['source_row']} has rule "
                f"{row['rule_text']!r}, which is neither a number nor one of "
                f"the action classes {sorted(ACTION_CLASSES)}.")
        if row["delta_logit_text"] != "N/A":
            raise SystemExit(
                f"{SHEET}: row {row['source_row']} carries the action class "
                f"{row['rule_text']!r} AND a shift of "
                f"{row['delta_logit_text']!r}. An action class is not an "
                "absorption effect and must not have one.")

    # THE TWO ROWS THAT SAY SO EXPLICITLY. Chelation reduces absorption of the
    # DRUG, and the sheet spells that out because the opposite reading would
    # cut a nutrient target for the wrong reason.
    timing = [r for r in interactions if r["rule_text"] == "TIMING"]
    if len(timing) != 2:
        raise SystemExit(
            f"{SHEET}: {len(timing)} rows are TIMING, expected two "
            "(Levothyroxine and Fluoroquinolones).")
    for row in timing:
        target = row["production_target"].lower()
        if "f_abs" not in target or not ("do not alter" in target
                                         or "unchanged" in target):
            raise SystemExit(
                f"{SHEET}: {row['drug']!r}'s production target no longer says "
                f"the nutrient's F_abs is untouched: "
                f"{row['production_target']!r}. That sentence is what stops a "
                "timing interaction being read as an absorption one.")

    # Statins x CoQ10 makes the same distinction from the other side.
    statins = next((r for r in interactions if r["drug"] == "Statins"), None)
    if statins is None or "not intestinal F_abs" not in statins["mechanism_class"]:
        raise SystemExit(
            f"{SHEET}: the Statins row no longer says its mechanism is not "
            "intestinal absorption. It is CRITICAL and it is not an F_abs "
            "effect, and the distinction is the point.")

    # THE SEVERITY SCALES DO NOT MATCH. The VETO registry's largest tier is
    # HIGH -- 111 of its 339 rows, just under a third -- and this sheet has no
    # HIGH at all, so an O9 row cannot express its biggest slice.
    veto = json.loads(
        (DATA_DIR / "veto_drug_nutrient_339.json").read_text(encoding="utf-8"))
    veto_rows = veto if isinstance(veto, list) else next(
        v for v in veto.values() if isinstance(v, list))
    veto_scale = {r["severity"] for r in veto_rows}
    if "HIGH" not in veto_scale:
        raise SystemExit(
            f"{SHEET}: the VETO registry no longer uses HIGH; the scale "
            "mismatch was recorded against it.")
    if "HIGH" in set(data["severity_scale"]):
        raise SystemExit(
            f"{SHEET}: this sheet now uses HIGH, so the two scales may agree. "
            "Re-read docs/parameter-gaps.md.")

    # And the disagreement that matters.
    if not data["severity_disagreements"]:
        raise SystemExit(
            f"{SHEET}: this sheet and the VETO registry now agree on every "
            "bridged pair. Re-read docs/parameter-gaps.md -- the insulin one "
            "was CRITICAL against MODERATE.")

    # WHAT THE INTERFACE OFFERS AND THIS SHEET DOES NOT MODEL, and the reverse.
    named = _ui_medications()
    if not named:
        raise SystemExit(f"{SHEET}: the UI contract names no Step 7 medication.")
    modelled = {r["drug"] for r in interactions}
    by_normalised = {_normalise(d) for d in modelled}
    unmodelled = sorted(m for m in named
                        if m not in modelled
                        and _normalise(m) not in by_normalised)
    if not unmodelled:
        raise SystemExit(
            f"{SHEET}: every medication the interface names now has a row. "
            "The finding is resolved -- re-read docs/parameter-gaps.md.")
    if "Diuretics" not in unmodelled:
        raise SystemExit(
            f"{SHEET}: 'Diuretics' now matches a row. It was the one UI name "
            "that is a BROADER class than the sheet's ('Thiazide/loop "
            "diuretics'), which is a different question from a drug with no "
            "row at all -- re-read docs/parameter-gaps.md.")
    if "Warfarin" not in modelled:
        raise SystemExit(
            f"{SHEET}: Warfarin is no longer modelled. It carried a CRITICAL "
            "Vitamin K veto and was the one drug the sheet models that the "
            "interface does not name.")
    if "Warfarin" in named:
        raise SystemExit(
            f"{SHEET}: the interface now names Warfarin, so the reverse gap "
            "is resolved. Re-read docs/parameter-gaps.md.")


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("usage: python -m sahacore.data.build_onboarding_o9 "
                         "<path-to-workbook>")
    data = extract(sys.argv[1])
    check(data)
    OUT.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n",
                   encoding="utf-8")

    interactions = data["interactions"]
    quantified = [r for r in interactions if r["legacy_multiplier"] is not None]
    critical = [r for r in interactions if r["severity"] == "CRITICAL"]
    named = _ui_medications()
    modelled = {r["drug"] for r in interactions}
    normalised = {_normalise(d) for d in modelled}
    spelled_differently = sorted(m for m in named
                                 if m not in modelled
                                 and _normalise(m) in normalised)
    unmodelled = sorted(m for m in named
                        if m not in modelled
                        and _normalise(m) not in normalised)
    print(f"wrote {OUT.name}")
    print(f"  {len(interactions)} drug-nutrient rows across {len(modelled)} drugs")
    print(f"  {len(quantified)} give a multiplier, and every Delta_logit_abs "
          "recomputes as ln(m) to 6 decimals -- O·O8's rule, verified")
    print(f"  {len(interactions) - len(quantified)} carry an action class "
          "instead (TIMING/MONITOR/VETO/CLEARANCE/N/A), none with a shift")
    print(f"  {len(critical)} CRITICAL")
    print(f"  {len(spelled_differently)} spelled differently but plainly the "
          "same drug: " + ", ".join(spelled_differently))
    print(f"  {len(unmodelled)} the interface names with NO row here: "
          + ", ".join(unmodelled))
    for d in data["severity_disagreements"]:
        print(f"  DISAGREES WITH THE VETO REGISTRY: {d['drug']} -- O9 "
              f"{d['o9_severity']}/{d['o9_rule']} against {d['veto_rule_id']} "
              f"{d['veto_severity']} ({d['veto_action'][:38]})")


if __name__ == "__main__":
    main()

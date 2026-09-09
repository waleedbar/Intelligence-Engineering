"""Extracts the engineer handoff's canonical invariants into
runtime_invariants.json.

Source: v39sEng2.xlsx, sheet '00_ENGINEER_START' -- "SahaPlusAI v39w —
Thin-Client Signal-Contract Engineer Handoff", described in its own row 4 as
the "Current engineer source of truth". Header at row 6, invariants at rows
7-29. Columns: Canonical invariant | Value | Engineering meaning | Backend
owner | Frontend owner | Data server owner | Gate.

    python -m sahacore.data.build_runtime_invariants <path-to-workbook>

WHY THIS IS FOUNDATION RATHER THAN CONFIGURATION. Most of these rows are not
settings anyone may choose; they are fail-closed safety positions, and the
sheet says so in its own words:

    Layer T          OFF                     "research-only until
                                              block-bootstrap and held-out
                                              superiority gates pass"
    Open Ear         BLOCKED                 "No audio/raw mic persistence"
    D14/D15          GLOBAL_MODIFIER_PENDING "No invented organ weights. Fail
                                              closed until evidence-locked
                                              mapping is signed off."
    Cluster coupling OFF (ETA=0)             "Numeric Gamma edges remain
                                              hypotheses"
    Actions          127 / 126 activatable   "Do not force 127 active arms."

A layer written without these in front of it can violate one by being
reasonable -- switching on a coupling term because the matrix is there, or
activating the 127th arm because the count looks short. Loading them before
the layers means the position is in the database rather than in whoever
remembers the sheet.

THE GATE COLUMN IS THE MACHINE-READABLE HALF. 18 of the 23 rows carry a gate
name (`nutrient_core_81`, `action_127_fail_closed`, `eta_net_zero`, ...); the
other five -- the behavioural sidecar rules -- carry none, and are kept with
a null gate rather than given one. Inventing a gate name for a row the sheet
left blank would make an unenforceable rule look enforced.

`enforced_by` is this build's own column, not the sheet's, in the same sense
as `resolved_by` in the parameter registry: it names the test that actually
executes the gate against what this repo contains, or is null when the gate
is recorded and not yet executable. It is filled in by hand below, one row at
a time, and tests/test_runtime_invariants.py asserts every name it mentions
is a test that exists.
"""
import json
import sys
from pathlib import Path

import openpyxl

SHEET = "00_ENGINEER_START"
HEADER_ROW = 6
OUT = Path(__file__).parent / "runtime_invariants.json"

COLUMNS = ["invariant", "value", "engineering_meaning", "backend_owner",
           "frontend_owner", "data_server_owner", "gate"]
HEADER_LABELS = ["Canonical invariant", "Value", "Engineering meaning",
                 "Backend owner", "Frontend owner", "Data server owner", "Gate"]

# Which gates this build can execute today, and the test that does it.
#
# A gate absent from this map is recorded and NOT yet enforced -- which is the
# honest state for a position about a layer nobody has written. It is not a
# licence to ignore it: the position still has to hold when that layer is
# written, and the gate is in the database to be read then.
ENFORCED_BY = {
    "schema_state_count_219":
        "test_runtime_invariants.py::test_gate_schema_state_count_219",
    "rb_partition_164_55_111":
        "test_runtime_invariants.py::test_gate_rb_partition_164_55",
    "nutrient_core_81":
        "test_runtime_invariants.py::test_gate_nutrient_core_81",
    "action_127_fail_closed":
        "test_runtime_invariants.py::test_gate_action_127_fail_closed",
    "veto_339_fk_coverage":
        "test_runtime_invariants.py::test_gate_veto_339",
    "eta_net_zero":
        "test_runtime_invariants.py::test_gate_eta_net_zero",
}

# The sheet's own count, checked against the extraction rather than the
# extraction being checked against itself.
EXPECTED_ROWS = 23
EXPECTED_GATED = 18


def _cell(ws, row: int, col: int) -> str | None:
    value = ws.cell(row=row, column=col).value
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def extract(workbook_path: str) -> list[dict]:
    ws = openpyxl.load_workbook(workbook_path, data_only=True)[SHEET]

    header = [_cell(ws, HEADER_ROW, c) for c in range(1, len(COLUMNS) + 1)]
    if header != HEADER_LABELS:
        raise SystemExit(f"{SHEET} header changed: expected {HEADER_LABELS}, got {header}")

    rows = []
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        if _cell(ws, r, 1) is None:
            continue
        row = {"source_row": r}
        for i, name in enumerate(COLUMNS, start=1):
            row[name] = _cell(ws, r, i)
        row["enforced_by"] = ENFORCED_BY.get(row["gate"]) if row["gate"] else None
        rows.append(row)
    return rows


def check(rows: list[dict]) -> None:
    if len(rows) != EXPECTED_ROWS:
        raise SystemExit(f"expected {EXPECTED_ROWS} invariants, extracted {len(rows)}")

    gated = [r for r in rows if r["gate"]]
    if len(gated) != EXPECTED_GATED:
        raise SystemExit(f"expected {EXPECTED_GATED} gated rows, got {len(gated)}")

    names = [r["gate"] for r in gated]
    if len(set(names)) != len(names):
        raise SystemExit("gate names are not unique")

    unknown = set(ENFORCED_BY) - set(names)
    if unknown:
        raise SystemExit(
            f"ENFORCED_BY names gates the sheet does not have: {sorted(unknown)}")

    for r in rows:
        if not r["value"] or not r["engineering_meaning"]:
            raise SystemExit(f"row {r['source_row']} has no value or meaning")
        if r["enforced_by"] and not r["gate"]:
            raise SystemExit(f"row {r['source_row']} claims enforcement without a gate")


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("usage: python -m sahacore.data.build_runtime_invariants <workbook.xlsx>")

    rows = extract(sys.argv[1])
    check(rows)
    OUT.write_text(json.dumps(rows, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    enforced = [r for r in rows if r["enforced_by"]]
    print(f"wrote {len(rows)} canonical invariants to {OUT.name}")
    print(f"   with a gate name : {sum(1 for r in rows if r['gate'])}")
    print(f"   enforced here    : {len(enforced)}")
    for r in rows:
        mark = "RUN " if r["enforced_by"] else ("    " if r["gate"] else "    ")
        print(f"   {mark}{(r['gate'] or '(no gate)'):<28} {r['value'][:44]}")


if __name__ == "__main__":
    main()

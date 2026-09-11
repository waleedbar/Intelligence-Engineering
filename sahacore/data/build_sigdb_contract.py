"""Extracts the SIGDB governance contract into sigdb_contract.json.

Source: SahaPlusAI_MASTER_..._v39sEng2SIGDB.xlsx, contract SIG-DB-v1.1, as of
2026-08-27. Sheets 12_STATE_BLOCKS, 17_FIREWALL, 18_GAPS, 19_QA_GATES and
21_DECISIONS.

    python -m sahacore.data.build_sigdb_contract <path-to-SIGDB.xlsx>

WHAT THIS FILE IS, because it is not another spec sheet. The master workbook
says how the engine computes. SIGDB says what may LEAVE it. Its own header:

    "Backend/data-server implementation contract for SahaTwin, SahaPulse,
     SahaAtlas and SahaPlan"
    "API posture: default deny / explicit allowlist"

and it is an audit as well as a contract: 263 XLSX sheets, 20 PDF pages and
75,520 populated cells reviewed, 1,156 formulas checked.

WHY IT IS IMPORTED BEFORE THE FIRST READ ENDPOINT AND NOT AFTER. FW01 is
DEFAULT_DENY -- "only explicitly registered API_OUTPUT fields may be
serialised". A default-deny posture is cheap to build into an API that has no
read endpoints yet, which is exactly what this service is today, and
expensive to retrofit onto one that already answers. Getting it wrong is not
a tidiness problem: FW05 forbids the 219-state vector, its covariance and the
raw nutrient pools from crossing the client boundary at all.

WHAT THE FIVE SHEETS ARE

    12_STATE_BLOCKS   the canonical 219-state partition, with each block's
                      Rao-Blackwell class and api_exposure
    17_FIREWALL       17 non-negotiable controls, FW01-FW17, each with the
                      test that proves it
    18_GAPS           68 audited findings with severity, evidence cells,
                      decision and release gate
    19_QA_GATES       34 release gates, most carrying an expected NUMBER --
                      which is why several can be enforced against registries
                      this build already loads
    21_DECISIONS      16 adjudications marked LOCKED_FOR_THIS_CONTRACT

NOTHING HERE IS THIS BUILD'S OPINION. The decisions are locked by their own
sheet and the gaps carry the auditor's verdicts; both are transcribed, and
where a gap contradicts something this build did, the gap wins and is
reported rather than argued with.
"""
import json
import sys
from pathlib import Path

import openpyxl

from sahacore.data.sheet_header import check_header

OUT = Path(__file__).parent / "sigdb_contract.json"

# sheet -> (header row, column labels, json key, expected row count)
TABLES = {
    "12_STATE_BLOCKS": (
        5,
        ["block_id", "block_name", "start_index", "end_index", "state_count",
         "rb_partition", "api_exposure", "source_sheet", "notes"],
        "state_blocks", 10),
    "17_FIREWALL": (
        5,
        ["control_id", "control_class", "rule", "severity", "test", "source",
         "implementation_owner", "status"],
        "firewall_controls", 17),
    "18_GAPS": (
        5,
        ["finding_id", "product", "severity", "topic", "finding", "evidence",
         "decision", "release_gate", "source_file", "status"],
        "gaps", 68),
    "19_QA_GATES": (
        5,
        ["qa_id", "severity", "rule", "expected", "actual_or_formula",
         "excel_formula", "status_formula", "source", "status"],
        "qa_gates", 34),
    "21_DECISIONS": (
        5,
        ["decision_id", "topic", "finding", "canonical_decision", "rationale",
         "owner", "status", "effective_release"],
        "decisions", 16),
}

# The contract's own identity, read from 00_README rather than typed here, so
# a newer SIGDB cannot be imported under the old version number by accident.
README_FIELDS = ["Contract version", "Source authority", "As of",
                 "Public projection", "Server state", "API posture",
                 "Clinical status"]

# Read off 18_GAPS and 17_FIREWALL rather than assumed. BLOCKER is the
# auditor's own top tier; a severity outside this set means the sheet grew a
# level and the loader's CHECK constraint would reject it anyway.
SEVERITIES = {"BLOCKER", "CRITICAL", "HIGH", "MEDIUM", "LOW"}

# 19_QA_GATES gates that name a plain integer AND a registry this build
# already loads. Verified here, in the extractor, because a contract whose
# release gates disagree with the data is not a contract worth storing.
#
# Deliberately NOT every gate: most read SIGDB's own sheets (COUNTIF over
# 11_ENTITY_AXES), and checking those would only prove the spreadsheet agrees
# with itself.
ENFORCEABLE_GATES = {
    "QA004": ("state_vector_219.json", None, 219),
    "QA005": ("nutrients_81.json", None, 81),
    "QA010": ("action_space_127.json", "actions", 127),
}


def _cell(ws, row: int, col: int):
    value = ws.cell(row=row, column=col).value
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value).strip()
    return text or None


def extract(workbook_path: str) -> dict:
    wb = openpyxl.load_workbook(workbook_path, data_only=True)

    contract = {}
    readme = wb["00_README"]
    for row in range(1, readme.max_row + 1):
        label = _cell(readme, row, 1)
        if label in README_FIELDS:
            contract[label] = _cell(readme, row, 2)
    missing = [f for f in README_FIELDS if f not in contract]
    if missing:
        raise SystemExit(f"00_README no longer states {missing}")

    data = {"contract": contract}
    for sheet, (header_row, labels, key, _) in TABLES.items():
        ws = wb[sheet]
        check_header(ws, sheet, header_row, 1, labels)

        rows = []
        for r in range(header_row + 1, ws.max_row + 1):
            if _cell(ws, r, 1) is None:
                continue
            row = {"source_row": r}
            for offset, label in enumerate(labels):
                row[label] = _cell(ws, r, 1 + offset)
            rows.append(row)
        data[key] = rows
    return data


def check(data: dict) -> None:
    for sheet, (_, _, key, expected) in TABLES.items():
        rows = data[key]
        if len(rows) != expected:
            raise SystemExit(
                f"{sheet}: expected {expected} rows, got {len(rows)}")
        id_column = next(iter(rows[0]))  # source_row
        id_column = list(rows[0])[1]     # the sheet's own id column
        ids = [r[id_column] for r in rows]
        if len(set(ids)) != len(ids):
            raise SystemExit(f"{sheet}: {id_column} is not unique")

    if data["contract"]["Contract version"] != "SIG-DB-v1.1":
        raise SystemExit(
            f"this is contract {data['contract']['Contract version']!r}, not "
            "SIG-DB-v1.1. A newer contract may have changed a locked "
            "decision; read it before importing it.")

    for table in ("firewall_controls", "gaps"):
        for row in data[table]:
            if row["severity"] not in SEVERITIES:
                raise SystemExit(
                    f"{table}: unknown severity {row['severity']!r} in "
                    f"{list(row.values())[1]}")

    # Every firewall control is non-negotiable by the sheet's own framing --
    # "Non-negotiable controls for API serialisation" -- so a control that
    # arrives OPTIONAL means the sheet changed its posture.
    #
    # THREE CARRY NO STATUS AND NO OWNER, and they are the three that matter
    # most to this service. FW15 ENGINE_COMPUTE_ONLY, FW16 SERIALIZER_ONLY_API
    # and FW17 NO_DEVICE_MODEL_ARTIFACT are all BLOCKER, all sourced to the
    # "v39w thin-client ruling", and all three leave implementation_owner and
    # status empty where FW01-FW14 fill both. They were appended later and the
    # last two columns were not carried down.
    #
    # Stored as they are, with the omission visible, and pinned by name: a
    # fourth statusless control stops the import rather than joining them.
    # Reported in docs/parameter-gaps.md.
    STATUSLESS = {"FW15", "FW16", "FW17"}
    statusless = {r["control_id"] for r in data["firewall_controls"]
                  if r["status"] is None}
    if statusless != STATUSLESS:
        raise SystemExit(
            f"firewall controls with no status are {sorted(statusless)}, not "
            f"{sorted(STATUSLESS)}. A BLOCKER control with no owner and no "
            "status is one nobody has accepted; read it rather than loading it.")
    for row in data["firewall_controls"]:
        if row["control_id"] in STATUSLESS:
            if row["implementation_owner"] is not None:
                raise SystemExit(
                    f"{row['control_id']} now names an owner; the finding that "
                    "it had none is stale.")
            continue
        if row["status"] != "REQUIRED":
            raise SystemExit(
                f"{row['control_id']} is {row['status']!r}, not REQUIRED. "
                "17_FIREWALL calls these non-negotiable.")

    # LOCKED_FOR_THIS_CONTRACT is what makes a decision binding rather than a
    # note. If one is ever unlocked, this build should stop and read it.
    unlocked = [d["decision_id"] for d in data["decisions"]
                if d["status"] != "LOCKED_FOR_THIS_CONTRACT"]
    if unlocked:
        raise SystemExit(f"decisions no longer locked: {unlocked}")

    # --- the state partition must be a partition -------------------------
    blocks = sorted(data["state_blocks"], key=lambda b: b["start_index"])
    if blocks[0]["start_index"] != 1:
        raise SystemExit("the state partition does not start at 1")
    if blocks[-1]["end_index"] != 219:
        raise SystemExit(
            f"the state partition ends at {blocks[-1]['end_index']}, not 219")
    for earlier, later in zip(blocks, blocks[1:]):
        if later["start_index"] != earlier["end_index"] + 1:
            raise SystemExit(
                f"{earlier['block_id']} ends at {earlier['end_index']} and "
                f"{later['block_id']} starts at {later['start_index']} -- the "
                "219 states are not covered exactly once.")
    for block in blocks:
        span = block["end_index"] - block["start_index"] + 1
        if span != block["state_count"]:
            raise SystemExit(
                f"{block['block_id']} spans {span} indices and claims "
                f"{block['state_count']}")
    total = sum(b["state_count"] for b in blocks)
    if total != 219:
        raise SystemExit(f"the blocks sum to {total}, not 219")

    # FW05 forbids the state vector from crossing the client boundary, so no
    # block may be exposed. If one ever is, that is a contradiction inside
    # the contract and not something to load quietly.
    exposed = [b["block_id"] for b in blocks if b["api_exposure"] != "HIDDEN"]
    if exposed:
        raise SystemExit(
            f"state blocks marked exposed: {exposed}. FW05 says no 219-state "
            "vector, covariance or raw pool crosses the client boundary.")

    # --- the gates this build can actually answer ------------------------
    # Three ways, because the gate itself has two halves. `expected` is TEXT
    # -- the number a human wrote -- and `actual_or_formula` is the value the
    # workbook computed, several of them by Excel formula. Comparing only one
    # would miss a sheet that disagrees with itself.
    for qa_id, (filename, key, expected) in ENFORCEABLE_GATES.items():
        gate = next((g for g in data["qa_gates"] if g["qa_id"] == qa_id), None)
        if gate is None:
            raise SystemExit(f"19_QA_GATES no longer contains {qa_id}")
        if gate["expected"] != str(expected):
            raise SystemExit(
                f"{qa_id} expects {gate['expected']!r}, and this build was "
                f"written against {expected}. Read the gate before changing "
                "the number.")
        if gate["actual_or_formula"] != expected:
            raise SystemExit(
                f"{qa_id} expects {expected} and the sheet's own computed "
                f"value is {gate['actual_or_formula']!r}. SIGDB disagrees "
                "with itself; do not load it.")
        loaded = json.loads(
            (Path(__file__).parent / filename).read_text(encoding="utf-8"))
        rows = loaded[key] if key else loaded
        if len(rows) != expected:
            raise SystemExit(
                f"{qa_id} requires {expected} and {filename} holds "
                f"{len(rows)}. The contract and the registry disagree.")

    # QA011: 126 activatable of 127 registered -- the difference is action
    # 127's safety hold, which this build already stores.
    actions = json.loads(
        (Path(__file__).parent / "action_space_127.json").read_text(encoding="utf-8"))
    held = [a for a in actions["actions"] if a.get("activation_hold")]
    if len(actions["actions"]) - len(held) != 126:
        raise SystemExit(
            f"QA011 requires 126 activatable actions; the registry has "
            f"{len(actions['actions'])} with {len(held)} held.")


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit(
            "usage: python -m sahacore.data.build_sigdb_contract <SIGDB.xlsx>")

    data = extract(sys.argv[1])
    check(data)
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n",
                   encoding="utf-8")

    print(f"wrote {OUT.name}  ({data['contract']['Contract version']}, "
          f"as of {data['contract']['As of']})")
    for _, (_, _, key, _) in TABLES.items():
        print(f"   {len(data[key]):>3}  {key}")
    blockers = [g for g in data["gaps"] if g["severity"] in ("BLOCKER", "CRITICAL")]
    print(f"   {len(blockers):>3}  of those gaps are BLOCKER or CRITICAL")


if __name__ == "__main__":
    main()

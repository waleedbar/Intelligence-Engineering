"""Extracts '★ Supplement Registry' into supplement_registry.json.

Source: v39sEng2.xlsx, sheet '★ Supplement Registry'.
'01_IMPORT_MANIFEST' order 42, role REGISTRY, import YES, backend Yes.

    python -m sahacore.data.build_supplement_registry <path-to-workbook>

THE NULL-CAP RULE is why this registry is separate from the 81 nutrients,
and the sheet states it plainly:

    "Where randomised trials of the supplement form are null, the engine caps
     its modelled effect at zero regardless of how strong the dietary-pattern
     evidence is. This already applies to vitamin D (VITAL) and EPA+DHA
     (STRENGTH, VITAL, ASCEND) in the cardiovascular weights... A null cap is
     not an omission -- it is the finding."

and the reason for the separation:

    "Keeping one registry would let a strong dietary association silently
     license a supplement claim the trials do not support."

That is a claim the engine could make by accident. A vitamin D capsule adds
to vit_d_iu, the dietary pattern for vitamin D carries real cardiovascular
weight, and without a cap the supplement inherits that weight -- which the
randomised trials do not support. The cap has to be in the data before Layer
C reads a dose, so it belongs in the foundation.

THE MAPPING COLUMN IS CLASSIFIED, NOT PARSED INTO KEYS. Its cells read
"YES -> vit_d_iu", "NO -- supplement only", "Partial -> glycine, proline",
"YES -> several". `mapping_kind` is the first word, which the sheet writes
consistently. `maps_to_nutrients` holds only the canonical ids that actually
appear and actually exist in the 81-nutrient registry -- "several" yields an
empty list rather than a guess about which several.

ONE ROW IS INCOMPLETE IN THE SOURCE. Betaine (trimethylglycine / TMG) has a
name and nothing else: no cap, no evidence position, no cluster permission.
It is loaded with those fields null and reported, rather than dropped (which
would hide a supplement the sheet means to cover) or filled in (which would
invent a safety position). See docs/parameter-gaps.md.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

SHEET = "★ Supplement Registry"
HEADER_ROW = 7
FIRST_COL = 2
OUT = Path(__file__).parent / "supplement_registry.json"

COLUMNS = ["supplement", "maps_to_canonical", "effect_cap", "evidence_position",
           "cluster_effect_permitted"]
LABELS = ["Supplement", "Maps to canonical nutrient?", "Effect cap",
          "Evidence position", "Cluster effect permitted"]

EXPECTED_ROWS = 13
INCOMPLETE = "Betaine (trimethylglycine / TMG)"

# The rows whose cap is stated as NULL. Named so the rule cannot be softened
# by an edit without this list disagreeing.
NULL_CAPPED = {"CoQ10 (ubiquinone / ubiquinol)", "Vitamin D", "EPA + DHA", "Curcumin"}


def _cell(ws, row: int, col: int) -> str | None:
    value = ws.cell(row=row, column=col).value
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def extract(workbook_path: str) -> list[dict]:
    ws = openpyxl.load_workbook(workbook_path, data_only=True)[SHEET]

    got = [_cell(ws, HEADER_ROW, c) for c in range(FIRST_COL, FIRST_COL + len(LABELS))]
    if got != LABELS:
        raise SystemExit(f"{SHEET} header changed: expected {LABELS}, got {got}")

    known = {n["id"] for n in json.loads(
        (Path(__file__).parent / "nutrients_81.json").read_text(encoding="utf-8"))}

    rows = []
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        name = _cell(ws, r, FIRST_COL)
        if name is None or name.startswith("BUILD NOTE"):
            continue
        row = {"source_row": r}
        for i, col in enumerate(COLUMNS, start=FIRST_COL):
            row[col] = _cell(ws, r, i)

        cell = row["maps_to_canonical"] or ""
        first = cell.split()[0].upper().rstrip(",") if cell else ""
        row["mapping_kind"] = first if first in ("YES", "NO", "PARTIAL") else None
        # Only ids the nutrient registry actually has. 'several' and
        # 'supplement only' yield nothing, which is the honest reading.
        row["maps_to_nutrients"] = [
            tok for tok in re.findall(r"[a-z0-9_]+", cell) if tok in known]
        row["is_complete"] = all(row[c] for c in COLUMNS)
        rows.append(row)
    return rows


def check(rows: list[dict]) -> None:
    if len(rows) != EXPECTED_ROWS:
        raise SystemExit(f"expected {EXPECTED_ROWS} supplements, got {len(rows)}")

    names = [r["supplement"] for r in rows]
    if len(set(names)) != len(names):
        raise SystemExit("supplement names are not unique")

    incomplete = [r["supplement"] for r in rows if not r["is_complete"]]
    if incomplete != [INCOMPLETE]:
        raise SystemExit(
            f"the incomplete rows are {incomplete}, not exactly [{INCOMPLETE!r}]. "
            "A newly incomplete row is a gap in the source that must be "
            "reported, not absorbed.")

    for r in rows:
        if r["is_complete"] and r["mapping_kind"] is None:
            raise SystemExit(
                f"{r['supplement']}: mapping cell {r['maps_to_canonical']!r} does "
                "not start with YES, NO or Partial")

    capped = {r["supplement"] for r in rows
              if r["effect_cap"] and "NULL CAP" in r["effect_cap"].upper()}
    if capped != NULL_CAPPED:
        raise SystemExit(
            f"the NULL CAP set is {sorted(capped)}, not {sorted(NULL_CAPPED)}. "
            "Adding or removing a null cap is a safety change and must be "
            "deliberate.")


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit(
            "usage: python -m sahacore.data.build_supplement_registry <workbook.xlsx>")

    rows = extract(sys.argv[1])
    check(rows)
    OUT.write_text(json.dumps(rows, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    print(f"wrote {len(rows)} supplements to {OUT.name}")
    for r in rows:
        cap = r["effect_cap"] or "— INCOMPLETE IN SOURCE —"
        maps = ",".join(r["maps_to_nutrients"]) or "—"
        print(f"   {r['supplement'][:32]:<34}{str(r['mapping_kind'] or '?'):<9}"
              f"{maps[:26]:<28}{cap[:34]}")
    print(f"   NULL CAP: {len(NULL_CAPPED)}   incomplete rows: "
          f"{sum(1 for r in rows if not r['is_complete'])}")


if __name__ == "__main__":
    main()

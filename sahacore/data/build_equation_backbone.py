"""Extracts '★ Equation Backbone' into equation_backbone.json.

Source: v39sEng2.xlsx, sheet '★ Equation Backbone' -- "Single consolidated
wiring map: every equation, what it reads, what it feeds, and on which
plane it runs". '01_IMPORT_MANIFEST' order 124, role CORE_ENGINE, import YES,
backend Yes.

    python -m sahacore.data.build_equation_backbone <path-to-workbook>

58 equations from ONB (onboarding warm-start) through RB (Rao-Blackwellised
factorisation), each with its formula, its inputs, its outputs and its
cadence. Where 'EQ · Canonical Build Rows' maps equations to Python
functions, this maps them to each other -- it is the dependency graph.

WHAT IT SETTLED. Three things this build had recorded as missing parameters
turn out here to be computed quantities, which is a different kind of answer
and a better one:

  CL   B4: "CL = CL_renal + CL_hepatic", with B5 giving
       CL_renal = GFR*f_filtered*(1-f_reabsorbed) from lab eGFR and B6 giving
       CL_hepatic = Q_H*E_H. It was reported MISSING_FK on B-002/B-003. It is
       not a parameter anyone forgot to write down; it is a sum of two others.

  rho  C2: "rho=exp(-k*dt); g=-expm1(-k*dt)/k". docs/parameter-gaps.md had
  g    hypothesised exactly this relation between the persistence recursion
       and #125 k_Z,k, and deliberately refused to encode it without the
       workbook stating it. The workbook states it here.

Those corrections are recorded in docs/parameter-gaps.md rather than quietly
absorbed.

eq_id IS NOT UNIQUE. `C2` appears twice: row 23 is Layer C's exact
excess-damage integration, row 61 is the Cost-Benefit net-value re-rank.
Different layers, different formulas, the same short id. This is the third
sheet in this build where an apparently-natural key is not one -- after
'PARAM · Eq Param FK' (A-001, A-002, B-001) and the VETO library's source
ids -- so the sheet's own row number is the identity here too.

NOTHING IS PARSED. The formula, input and output cells are prose-and-notation
mixtures ("O6/O8/O11/O13 + Step-12 labs", "→ E (219-D state init), M0"). They
are stored verbatim. Turning the arrows into a graph would be a second,
larger job, and one this sheet does not authorise on its own.
"""
import json
import sys
from pathlib import Path

import openpyxl

SHEET = "★ Equation Backbone"
HEADER_ROW = 5
FIRST_COL = 2
OUT = Path(__file__).parent / "equation_backbone.json"

COLUMNS = ["eq_id", "layer", "name", "formula", "inputs", "outputs", "cadence"]
LABELS = ["ID", "Layer", "Name", "Formula", "Inputs ← (reads from)",
          "Outputs → (feeds)", "Cadence / execution plane"]

EXPECTED_ROWS = 58
# The one id the sheet reuses, and the two layers that share it.
DUPLICATE_ID = "C2"


def _cell(ws, row: int, col: int) -> str | None:
    value = ws.cell(row=row, column=col).value
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def extract(workbook_path: str) -> list[dict]:
    ws = openpyxl.load_workbook(workbook_path, data_only=True)[SHEET]

    got = [_cell(ws, HEADER_ROW, c) for c in range(FIRST_COL, FIRST_COL + len(LABELS))]
    if got != LABELS:
        raise SystemExit(f"{SHEET} header changed: expected {LABELS}, got {got}")

    rows = []
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        if _cell(ws, r, FIRST_COL) is None:
            continue
        row = {"source_row": r}
        for i, col in enumerate(COLUMNS, start=FIRST_COL):
            row[col] = _cell(ws, r, i)
        rows.append(row)
    return rows


def check(rows: list[dict]) -> None:
    if len(rows) != EXPECTED_ROWS:
        raise SystemExit(f"expected {EXPECTED_ROWS} equations, got {len(rows)}")

    for r in rows:
        if not r["layer"] or not r["name"] or not r["formula"]:
            raise SystemExit(
                f"row {r['source_row']} ({r['eq_id']}) is missing layer, name or formula")

    ids = [r["eq_id"] for r in rows]
    repeated = {i for i in ids if ids.count(i) > 1}
    if repeated != {DUPLICATE_ID}:
        raise SystemExit(
            f"the repeated ids are {sorted(repeated)}, not {{{DUPLICATE_ID!r}}}. "
            "eq_id is not the identity here -- source_row is -- but a NEW "
            "collision is worth looking at before it is absorbed.")
    layers = {r["layer"] for r in rows if r["eq_id"] == DUPLICATE_ID}
    if len(layers) != 2:
        raise SystemExit(f"the two {DUPLICATE_ID} rows should be in different layers: {layers}")


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit(
            "usage: python -m sahacore.data.build_equation_backbone <workbook.xlsx>")

    rows = extract(sys.argv[1])
    check(rows)
    OUT.write_text(json.dumps(rows, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    layers: dict[str, int] = {}
    for r in rows:
        layers[r["layer"]] = layers.get(r["layer"], 0) + 1
    print(f"wrote {len(rows)} equations to {OUT.name}")
    for layer, n in layers.items():
        print(f"   {n:>3}  {layer}")
    print(f"   ids reused across layers: {DUPLICATE_ID}")


if __name__ == "__main__":
    main()

"""Extracts '★ Validation Test Battery' into validation_battery.json.

Source: v39sEng2.xlsx, sheet '★ Validation Test Battery'.
'01_IMPORT_MANIFEST' order 14, role VALIDATION, import YES, backend Yes.

    python -m sahacore.data.build_validation_battery <path-to-workbook>

WHAT THE SHEET IS. The engine's acceptance criteria: named tests, each with
a target, a method, a pass criterion and a gate. Its own framing, verbatim:

    "★ VALIDATION & ADMISSION TEST BATTERY  ·  verification is not clinical
     validation"
    "Named mathematical, software and model-admission tests. Synthetic truth
     can verify code. It cannot validate physiology."

and how to read the levels:

    "Component and subsystem tests verify implementation properties.
     Integrative tests..."

WHY IT IS IMPORTED IN PHASE 1, before most of what it tests exists. Because
a criterion that arrives after the code is a criterion the code was not
written to meet. Several of these gates are already executable against
registries this build has loaded -- I7 and I8 are column sums over the damage
and nutrient-cluster registries, I16 is the VETO primary-key integrity check,
I6 is the firewall -- and importing the battery is what made it visible that
two of them were being asserted here at a LOOSER tolerance than the battery
requires. That is the value: the acceptance criteria stop living in a
spreadsheet nobody diffs against.

THE COVERAGE COLUMNS ARE THIS BUILD'S, NOT THE SHEET'S, in the same sense as
parameter_registry.resolved_by and runtime_invariant.enforced_by:

    ENFORCED  a named test in this repo executes the stated criterion.
    PARTIAL   part of the criterion is executed and part is not; the note
              says which part, and why the rest cannot run yet.
    NOT_YET   nothing in this repo executes it -- usually because the layer
              it tests has not been built.

Nothing is marked ENFORCED on resemblance. LM-P01 to LM-P03 are the clearest
case: 'Live Verification Lab' exercises all three, but at ONE point, and the
battery's method says "Sweep S, Z, theta, alpha, beta, dt". A worked example
is not a sweep, so they are PARTIAL.

ONE TEST WAS DROPPED BY THE SHEET and the note explaining why is extracted
with the rest, because a removed acceptance test is a decision, not an
absence.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

SHEET = "★ Validation Test Battery"
FIRST_COL = 2
OUT = Path(__file__).parent / "validation_battery.json"

COLUMNS = ["target", "method", "pass_criterion", "gate"]
HEADER_LABELS = ["ID", "Target", "Method", "Pass criterion", "Gate"]

# The battery's own id families. A row whose first cell matches one of these
# and which carries a target is a test; everything else is a section heading
# or prose.
TEST_ID = re.compile(r"^(?:C|S|I|SI|SV)\d+$|^LM-P\d+$|^T-R\d+$")

# The gates the sheet actually uses. A new one is a new admission policy, so
# the extractor stops rather than importing a word it has not seen.
KNOWN_GATES = {
    "BLOCKING",
    "MONITORED",
    "BLOCKING_FOR_PROMOTION",
    "BLOCKING_FOR_FLAG",
    "BLOCKING_FOR_SPECTRAL_CODE",
    "BLOCKING_FOR_SHADOW",
}

# Row 35 records a test the sheet decided not to keep.
DROPPED_NOTE_ROW = 35

# ---------------------------------------------------------------------------
# What this repo actually executes, keyed by test id.
#
# Each entry is (coverage, enforced_by, note). Written here rather than
# guessed by matching names, and `check()` refuses any id that is not in the
# battery -- so a test that is renamed upstream fails the extract instead of
# quietly losing its coverage.
COVERAGE: dict[str, tuple[str, str | None, str]] = {
    "I6": (
        "ENFORCED", "tests/test_firewall.py",
        "The T-1 firewall suite attempts engine reads through every "
        "client-facing path and asserts role-level denial, which is the "
        "criterion as written.",
    ),
    "I7": (
        "ENFORCED",
        "tests/test_layer_c_d_registries.py::test_each_cluster_weight_pct_sums_to_100",
        "All twelve columns sum to exactly 100.0. The assertion was written "
        "at abs=1e-2 before this sheet was imported and is now held to the "
        "battery's 'exactly'.",
    ),
    "I8": (
        "ENFORCED",
        "tests/test_layer_c_d_registries.py::test_each_cluster_column_sums_to_one",
        "Twelve columns, each 1.0 to within floating-point summation error. "
        "Was abs=1e-4, now held to the battery's 1.000000.",
    ),
    "I16": (
        "ENFORCED",
        "tests/test_veto_registry.py::test_veto_01_339_rows_and_339_unique_canonical_ids",
        "339 rows, 339 unique canonical ids, 319 distinct retained source "
        "ids. Also asserted in CI.",
    ),
    "C14": (
        "ENFORCED",
        "tests/test_bistability_guard.py::test_each_cluster_sits_exactly_on_its_own_cap",
        "Evaluated at the registry's declared maximum alpha/beta per "
        "cluster, where the product equals the cap by construction. A "
        "running engine's own alpha and beta would be checked by the same "
        "function, engine_internal.bistability_margin.",
    ),
    "C16": (
        "ENFORCED",
        "tests/test_validation_battery.py::test_c16_theta_elastic_is_ratified_per_cluster",
        "All twelve clusters carry a theta_elastic inside 42-70 AU.",
    ),
    "C17": (
        "PARTIAL",
        "tests/test_bistability_guard.py",
        "B_k = cap_k/(gamma_k*r_k) is implemented as "
        "engine_internal.bistability_margin and asserted. The second half, "
        "p_viol,k < 0.01, needs a posterior over (gamma, r) that Phase 2's "
        "cohort fit has not produced -- the sheet says so itself.",
    ),
    "C5": (
        "PARTIAL",
        "tests/test_verification_labs.py::test_zero_order_hold_gain_matches_the_published_grid",
        "The convention g = (1-e^(-k*dt))/k is verified and reproduced to "
        "1e-12. Comparing each discrete update against its own block's "
        "analytic solution needs the blocks, which are not built.",
    ),
    "C1": (
        "PARTIAL",
        "tests/test_verification_labs.py::test_gamma_kernel_integrates_to_one",
        "The kernel's normalisation is verified against the workbook's "
        "published integral, but on the demo grid, whose residual is 5.2e-3 "
        "-- four orders above this criterion. Reaching 1e-6 by trapezoid "
        "needs dt near 0.01 min, because the rule converges at O(dt^k) in "
        "the smallest shape parameter, not O(dt^2). See docs/parameter-gaps.md.",
    ),
    "LM-P01": (
        "PARTIAL", "tests/test_verification_labs.py::test_lm_p01_bounds_hold",
        "Verified at the workbook's single worked point. The method asks for "
        "a sweep over S, Z, theta, alpha, beta and dt.",
    ),
    "LM-P02": (
        "PARTIAL",
        "tests/test_verification_labs.py::test_lm_p02_two_half_steps_equal_one_full_step",
        "Two half-steps are checked against one full step. The method asks "
        "for one day against 24 hourly steps.",
    ),
    "LM-P03": (
        "PARTIAL",
        "tests/test_verification_labs.py::test_layer_m_scar_update_reproduces_every_published_quantity",
        "tau = 1/beta and t-half = ln(2)/beta are verified at the worked "
        "point.",
    ),
    "C8": (
        "PARTIAL", "tests/test_verification_labs.py::test_this_build_keeps_eta_net_at_zero",
        "'eta remains 0' is asserted against the registry and the runtime "
        "invariant. The rest of the criterion -- held-out calibration and "
        "max Re eig(J_full) -- needs a fitted model and the full Jacobian.",
    ),
}


def _text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _cell(ws, row: int, offset: int):
    return ws.cell(row=row, column=FIRST_COL + offset).value


def extract(path: str) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET]
    rows = list(ws.iter_rows(max_col=FIRST_COL + len(COLUMNS), values_only=True))

    tests: list[dict] = []
    sections: list[dict] = []
    section: str | None = None
    headers_seen = 0

    for number, raw in enumerate(rows, 1):
        first = _text(raw[FIRST_COL - 1])
        if first is None:
            continue
        second = _text(raw[FIRST_COL])

        if first == "ID":
            # A header row. Confirm the columns are where they are expected.
            found = [_text(_cell(ws, number, i)) for i in range(len(HEADER_LABELS))]
            if found != HEADER_LABELS:
                raise SystemExit(
                    f"{SHEET}: header row {number} reads {found}, expected "
                    f"{HEADER_LABELS}. The sheet moved.")
            headers_seen += 1
            continue

        if second is None:
            section = first
            sections.append({"source_row": number, "heading": first})
            continue

        if not TEST_ID.match(first):
            # Prose in the first column with something beside it. The sheet
            # has none today; refusing keeps a silently-dropped test loud.
            raise SystemExit(
                f"{SHEET}: row {number} has {first!r} in the id column with "
                f"{second[:60]!r} beside it, and {first!r} is not a test id. "
                "Either the id families changed or a test was reformatted.")

        tests.append({
            "test_id": first,
            "source_row": number,
            "section": section,
            **{name: _text(_cell(ws, number, i + 1))
               for i, name in enumerate(COLUMNS)},
        })

    dropped = _text(_cell(ws, DROPPED_NOTE_ROW, 0))
    wb.close()

    for test in tests:
        coverage, enforced_by, note = COVERAGE.get(
            test["test_id"], ("NOT_YET", None, ""))
        test["coverage"] = coverage
        test["enforced_by"] = enforced_by
        test["coverage_note"] = note or None

    return {
        "sheet": SHEET,
        "headers_seen": headers_seen,
        "sections": sections,
        "dropped_test_note": dropped,
        "tests": tests,
    }


def check(data: dict) -> None:
    tests = data["tests"]
    ids = [t["test_id"] for t in tests]
    if len(ids) != len(set(ids)):
        duplicates = sorted({i for i in ids if ids.count(i) > 1})
        raise SystemExit(f"{SHEET}: duplicate test ids {duplicates}.")

    for test in tests:
        for field in ("target", "method", "pass_criterion", "gate"):
            if not test[field]:
                raise SystemExit(
                    f"{SHEET}: {test['test_id']} at row {test['source_row']} "
                    f"has no {field}. A test without one of those is not "
                    "admissible as an acceptance criterion.")
        if test["gate"] not in KNOWN_GATES:
            raise SystemExit(
                f"{SHEET}: {test['test_id']} carries gate {test['gate']!r}, "
                f"which is not one of {sorted(KNOWN_GATES)}. A new gate is a "
                "new admission policy -- read it before importing it.")
        if test["section"] is None:
            raise SystemExit(
                f"{SHEET}: {test['test_id']} appears before any section "
                "heading.")

    unknown = sorted(set(COVERAGE) - set(ids))
    if unknown:
        raise SystemExit(
            f"{SHEET}: coverage is declared for {unknown}, which the battery "
            "does not contain. A renamed test must not silently keep its "
            "coverage.")

    for test in tests:
        if test["coverage"] == "NOT_YET":
            if test["enforced_by"] or test["coverage_note"]:
                raise SystemExit(
                    f"{test['test_id']}: NOT_YET cannot name a test or a note.")
        elif not test["enforced_by"] or not test["coverage_note"]:
            raise SystemExit(
                f"{test['test_id']}: {test['coverage']} must name the test "
                "that runs it and say what it covers.")

    if not data["dropped_test_note"] or "DROPPED" not in data["dropped_test_note"]:
        raise SystemExit(
            f"{SHEET}: row {DROPPED_NOTE_ROW} no longer carries the note about "
            "the dropped test.")

    if data["headers_seen"] < 4:
        raise SystemExit(
            f"{SHEET}: only {data['headers_seen']} header rows were confirmed; "
            "the column check is not actually running over the sheet.")


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("usage: python -m sahacore.data.build_validation_battery "
                         "<path-to-workbook>")
    data = extract(sys.argv[1])
    check(data)
    OUT.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n",
                   encoding="utf-8")

    gates: dict[str, int] = {}
    coverage: dict[str, int] = {}
    for test in data["tests"]:
        gates[test["gate"]] = gates.get(test["gate"], 0) + 1
        coverage[test["coverage"]] = coverage.get(test["coverage"], 0) + 1
    print(f"wrote {OUT.name}")
    print(f"  {len(data['tests'])} tests over {len(data['sections'])} sections")
    print("  gates:    " + ", ".join(f"{k} {v}" for k, v in sorted(gates.items())))
    print("  coverage: " + ", ".join(f"{k} {v}" for k, v in sorted(coverage.items())))


if __name__ == "__main__":
    main()

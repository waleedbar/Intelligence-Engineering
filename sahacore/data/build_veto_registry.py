"""Extracts the drug-nutrient VETO library into veto_drug_nutrient_339.json.

Source: v39sEng2.xlsx, sheet 'MERGE·VETO Drug-Nutrient 339' -- "★ v37.1 —
DRUG × NUTRIENT VETO LIBRARY (339 rules) · engine safety data · canonical IDs
repaired". Columns: rule_id | drug_or_class | nutrient_or_food |
engine_nutrient_id | nutrient_category | link_type | severity | action |
bandit_action | message_id | clinical_rationale | k_ij | source_rule_id.

    python -m sahacore.data.build_veto_registry <path-to-workbook>

WHICH SHEET IS AUTHORITATIVE. The workbook also has a sheet literally named
'VETO Canonical 339'. It is NOT the one to load, and says so in its own third
row:

    "REFERENCE ONLY -- the sheet name is historical and does not guarantee
     the active registry row count. Production/build loaders MUST use
     MERGE·VETO Drug-Nutrient 339, the active registry with 339 unique
     canonical rule_ids."

It holds 266 rows. Loading it because its name matches would have shipped a
truncated safety table, so this loader reads the sheet the workbook names and
tests/test_veto_registry.py pins that choice.

WHY THE TWO ID COLUMNS. 'Replay Contract' section E records the defect and
its repair: "339 rows contained only 319 unique rule IDs. A unique
VETO-DN-0001…0339 rule_id is canonical; original duplicate source IDs are
preserved separately." So `rule_id` is the primary key and `source_rule_id`
is kept verbatim, duplicates and all -- 339 of the former, 319 distinct
values of the latter. Neither is discarded to make the other look tidy.

NOTHING IS DERIVED HERE. Every field is the cell, transcribed. The two
invariants the schema enforces were read off the data rather than imposed on
it, and are checked here before writing:

  * engine_nutrient_id is present exactly when link_type is 'engine_state'
    (250 rows). The other 89 -- supplement_offmodel, food_flag, drug_drug,
    review, nutrient_group -- name no engine state, and an id invented for
    them would be a foreign key to a nutrient the rule is not about.
  * severity determines bandit_action, with no exceptions in 339 rows.
    CRITICAL removes the arm; it never merely warns. This is the one that
    matters if a later edit is careless, so the database refuses the row.
"""
import json
import sys
from pathlib import Path

import openpyxl

from sahacore.data.sheet_header import check_header

SHEET = "MERGE·VETO Drug-Nutrient 339"
REFERENCE_ONLY_SHEET = "VETO Canonical 339"
HEADER_ROW = 3
OUT = Path(__file__).parent / "veto_drug_nutrient_339.json"

COLUMNS = [
    "rule_id", "drug_or_class", "nutrient_or_food", "engine_nutrient_id",
    "nutrient_category", "link_type", "severity", "action", "bandit_action",
    "message_id", "clinical_rationale", "k_ij", "source_rule_id",
]

# Read off the 339 rows, not imposed on them. Asserted below and enforced by
# a CHECK constraint in sql/014_veto_action_space.sql.
SEVERITY_TO_BANDIT_ACTION = {
    "CRITICAL": "HARD_VETO (remove arm)",
    "HIGH": "SOFT_PENALTY (down-weight)",
    "MODERATE": "WARNING (show note)",
    "CONTROVERSIAL": "WARNING (show note)",
    "LOW": "INFORMATIONAL (passive)",
}

# The sheet's own header states the distribution, so the extraction can be
# checked against the source's own arithmetic rather than against itself.
DECLARED_SEVERITY_COUNTS = {
    "CRITICAL": 94, "HIGH": 111, "MODERATE": 110, "LOW": 23, "CONTROVERSIAL": 1,
}

# link_type values whose rules name an engine state. Only these carry an
# engine_nutrient_id.
ENGINE_STATE_LINK = "engine_state"


def _cell(ws, row: int, col: int) -> str | None:
    value = ws.cell(row=row, column=col).value
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def extract(workbook_path: str) -> list[dict]:
    wb = openpyxl.load_workbook(workbook_path, data_only=True)

    check_header(wb[SHEET], SHEET, HEADER_ROW, 1, COLUMNS)

    ws = wb[SHEET]
    rows = []
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        if _cell(ws, r, 1) is None:
            continue
        row = {"source_row": r}
        for i, name in enumerate(COLUMNS, start=1):
            row[name] = _cell(ws, r, i)
        row["k_ij"] = int(row["k_ij"])
        rows.append(row)
    return rows


def check(rows: list[dict]) -> None:
    """Refuse to write a file that would not survive its own schema."""
    if len(rows) != 339:
        raise SystemExit(f"expected 339 rules, extracted {len(rows)}")

    rule_ids = [r["rule_id"] for r in rows]
    if len(set(rule_ids)) != 339:
        raise SystemExit("rule_id is not unique -- it is the canonical primary key")
    expected = [f"VETO-DN-{n:04d}" for n in range(1, 340)]
    if sorted(rule_ids) != expected:
        raise SystemExit("rule_ids are not the contiguous VETO-DN-0001..0339 range")

    counts: dict[str, int] = {}
    for r in rows:
        counts[r["severity"]] = counts.get(r["severity"], 0) + 1
    if counts != DECLARED_SEVERITY_COUNTS:
        raise SystemExit(
            f"severity distribution {counts} does not match the sheet's own "
            f"header, which states {DECLARED_SEVERITY_COUNTS}")

    for r in rows:
        want = SEVERITY_TO_BANDIT_ACTION.get(r["severity"])
        if r["bandit_action"] != want:
            raise SystemExit(
                f"{r['rule_id']}: severity {r['severity']} carries bandit_action "
                f"{r['bandit_action']!r}, not {want!r}")
        has_id = r["engine_nutrient_id"] is not None
        if has_id != (r["link_type"] == ENGINE_STATE_LINK):
            raise SystemExit(
                f"{r['rule_id']}: link_type {r['link_type']!r} with "
                f"engine_nutrient_id {r['engine_nutrient_id']!r}")

    known = {n["id"] for n in json.loads(
        (Path(__file__).parent / "nutrients_81.json").read_text(encoding="utf-8"))}
    unknown = {r["engine_nutrient_id"] for r in rows
               if r["engine_nutrient_id"] and r["engine_nutrient_id"] not in known}
    if unknown:
        raise SystemExit(f"engine_nutrient_ids absent from the 81-nutrient registry: {unknown}")


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("usage: python -m sahacore.data.build_veto_registry <workbook.xlsx>")

    rows = extract(sys.argv[1])
    check(rows)
    OUT.write_text(json.dumps(rows, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    sources = [r["source_rule_id"] for r in rows]
    print(f"wrote {len(rows)} VETO rules to {OUT.name}")
    print(f"   canonical rule_id unique      : {len(set(r['rule_id'] for r in rows))}")
    print(f"   source_rule_id distinct       : {len(set(sources))} "
          f"({len(sources) - len(set(sources))} preserved duplicates)")
    print(f"   carrying an engine_nutrient_id: "
          f"{sum(1 for r in rows if r['engine_nutrient_id'])}")
    for sev, n in sorted(DECLARED_SEVERITY_COUNTS.items()):
        print(f"   {sev:<14}{n:>4}  -> {SEVERITY_TO_BANDIT_ACTION[sev]}")


if __name__ == "__main__":
    main()
